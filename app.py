"""
EV Battery Supplier RFP Evaluator  (Streamlit)
------------------------------------------------
Validates supplier responses against the "EV Traction Battery Supplier RFP and
Evaluation Template v1.0": 8 mandatory pass/fail gates + 9 weighted criteria
(100 points, 0-5 evidence score each) + decision thresholds.

Run:   pip install -r requirements.txt
       streamlit run app.py
"""
import io
import json
import os
import re
import sqlite3
import time
import datetime as dt

import pandas as pd
import streamlit as st

try:
    import anthropic
except ImportError:  # app still works in offline-heuristic mode
    anthropic = None

DB_PATH = os.environ.get("RFP_DB", "rfp_evaluator.db")

# ----------------------------------------------------------------------------
# RFP TEMPLATE RUBRIC (taken from the attached template)
# ----------------------------------------------------------------------------
GATES = [
    ("Safety and legal compliance", "Market compliance plan; certification/test status; deviations register; named accountable owner",
     ["compliance", "ece r100", "r100", "un reg", "certif", "deviation", "accountable", "homologation", "type-approval", "type approval"]),
    ("Quality management", "IATF 16949 production-site certification or buyer-approved transition plan; ISO 9001 minimum",
     ["iatf", "16949", "iso 9001", "quality management"]),
    ("No leakage / containment", "Sealing strategy, ingress target, test plan, venting/drain concept, coolant and electrolyte containment evidence",
     ["leak", "seal", "ingress", "ip67", "ip6k9k", "helium", "containment", "venting"]),
    ("Thermal safety", "Thermal-propagation and abuse-test strategy; detection, mitigation, emergency response, test evidence",
     ["thermal runaway", "propagation", "abuse", "thermal safety", "mitigation", "fire", "gas detection"]),
    ("Traceability", "Cell-to-pack genealogy, software/config control, material-batch traceability, recall support",
     ["traceab", "genealogy", "batch", "serial", "recall", "lot "]),
    ("Cybersecurity", "BMS and vehicle-interface security plan, vulnerability management, incident notification, data controls",
     ["cyber", "21434", "secure boot", "vulnerab", "penetration", "sbom", "encryption", "authentication"]),
    ("Responsible end of life", "Collection, transport, repair/reuse/recycling plan; downstream partner qualifications; material recovery evidence",
     ["recycl", "end of life", "end-of-life", "take-back", "second-life", "second life", "reuse", "recovery"]),
    ("Financial / continuity", "Financial viability, site capacity, business-continuity and sub-tier risk plan",
     ["financial", "continuity", "bcp", "capacity", "sub-tier", "audited", "revenue"]),
]

CRITERIA = [
    ("design", "1. Design, packaging & serviceability", 8,
     "Efficient, vehicle-integrated, crash-aware, repairable design with controlled interfaces and credible mass/volume data.",
     ["envelope", "mass", "cad", "mounting", "crash", "underbody", "serviceab", "hvil", "service disconnect", "dfm", "tolerance"]),
    ("leakage", "2. Leakage / sealing & environmental robustness", 15,
     "Validated ingress and coolant/electrolyte containment strategy; production leak testing; robust post-test evidence.",
     ["leak", "seal", "ingress", "ip67", "helium", "pressure decay", "end-of-line", "corrosion", "humidity"]),
    ("vibration", "3. Vibration, shock & mechanical durability", 12,
     "Applicable road-load / vibration / shock results; retention, fatigue, isolation and post-test performance evidence.",
     ["vibration", "shock", "road-load", "road load", "fatigue", "durability", "mechanical", "insulation resistance"]),
    ("thermal", "4. Thermal management & thermal-cycle safety", 15,
     "Validated thermal model and test results; temperature uniformity; thermal-cycle / propagation / mitigation evidence.",
     ["thermal", "cooling", "temperature uniform", "thermal cycl", "propagation", "runaway", "sensor", "preconditioning"]),
    ("materials", "5. Materials, traceability & responsible sourcing", 12,
     "Full material and critical-mineral traceability; compliance, lifecycle-carbon, recycled-content and supplier controls.",
     ["cathode", "anode", "electrolyte", "cobalt", "lithium", "nickel", "imds", "due diligence", "carbon footprint", "recycled content", "provenance"]),
    ("performance", "6. Charging, capacity & lifecycle performance", 15,
     "Measured capacity, fast-charge, power, degradation, SOH and warranty performance under stated conditions.",
     ["kwh", "capacity", "fast-charge", "fast charge", "c-rate", "cycle life", "calendar", "degradation", "soh", "warranty", "internal resistance"]),
    ("eol", "7. Disposal, repair, reuse & recycling", 10,
     "Practical take-back, repair, second-life and recycling plan with qualified partners and evidence of recovery pathway.",
     ["take-back", "recycl", "second-life", "second life", "remanufactur", "repair", "disassembl", "recovery"]),
    ("security", "8. Security, functional safety & software governance", 8,
     "Credible cybersecurity and functional-safety lifecycle, secure updates, test evidence, disclosure and response process.",
     ["26262", "21434", "asil", "secure boot", "ota", "update", "hazard analysis", "diagnostic coverage", "incident"]),
    ("commercial", "9. Commercial, capacity & supplier resilience", 5,
     "Transparent TCO; viable capacity; PPAP/APQP plan; BCP, financial health, logistics and warranty commitments.",
     ["tco", "price", "ppap", "apqp", "bcp", "capacity", "logistics", "financial", "warranty"]),
]
CRIT_KEYS = [c[0] for c in CRITERIA]
SCORE_ANCHORS = {5: "Leading", 4: "Strong", 3: "Acceptable", 2: "Weak", 1: "Poor", 0: "No evidence"}

THRESHOLDS = [(85, "Preferred"), (75, "Competitive"), (65, "Conditional"), (0, "Do not progress")]


def band(total, gates_ok=True):
    if not gates_ok:
        return "Gate failed - not ranked"
    for lim, name in THRESHOLDS:
        if total >= lim:
            return name
    return "Do not progress"


def stars(total):
    return round(total / 20, 1)  # 0-5 star equivalent


def weighted_total(scores):
    return round(sum((scores.get(k, {}).get("score", 0) / 5) * w for k, _, w, _, _ in CRITERIA), 1)


# ----------------------------------------------------------------------------
# DATABASE  (memory of runs + audit history)
# ----------------------------------------------------------------------------
SCHEMA = """
CREATE TABLE IF NOT EXISTS suppliers(id INTEGER PRIMARY KEY, name TEXT UNIQUE, contact TEXT, created_at TEXT);
CREATE TABLE IF NOT EXISTS responses(id INTEGER PRIMARY KEY, supplier_id INTEGER, rfp_ref TEXT, title TEXT,
    text TEXT, cost REAL, currency TEXT, filename TEXT, created_at TEXT);
CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY, response_id INTEGER, mode TEXT, model TEXT, started TEXT,
    duration_s REAL, total_score REAL, rating TEXT, gates_passed INTEGER, gates_json TEXT, scores_json TEXT,
    summary TEXT, tokens_in INTEGER, tokens_out INTEGER, status TEXT, error TEXT);
CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY, ts TEXT, entity TEXT, entity_id TEXT, action TEXT, detail TEXT);
CREATE TABLE IF NOT EXISTS company_batteries(id INTEGER PRIMARY KEY, chemistry TEXT, application TEXT, notes TEXT);
CREATE TABLE IF NOT EXISTS supplier_batteries(id INTEGER PRIMARY KEY, supplier TEXT, chemistry TEXT, capacity_kwh REAL,
    voltage_v REAL, price_per_kwh REAL, source TEXT);
CREATE TABLE IF NOT EXISTS supplier_profiles(supplier TEXT PRIMARY KEY, years_in_business INTEGER, hq TEXT, specialties TEXT,
    certifications TEXT, key_customers TEXT, expertise_level INTEGER, notes TEXT, updated TEXT);
CREATE TABLE IF NOT EXISTS materials(id INTEGER PRIMARY KEY, supplier TEXT, material TEXT, category TEXT, origin TEXT,
    recycled_pct REAL, source TEXT);
"""


def conn():
    c = sqlite3.connect(DB_PATH, check_same_thread=False)
    c.row_factory = sqlite3.Row
    return c


def init_db():
    with conn() as c:
        c.executescript(SCHEMA)


def now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def execute(sql, params=()):
    with conn() as c:
        cur = c.execute(sql, params)
        return cur.lastrowid


def query(sql, params=()):
    with conn() as c:
        return pd.read_sql_query(sql, c, params=params)


def log(entity, entity_id, action, detail=""):
    execute("INSERT INTO audit(ts,entity,entity_id,action,detail) VALUES(?,?,?,?,?)",
            (now(), entity, str(entity_id), action, detail[:4000]))


def get_or_create_supplier(name, contact=""):
    name = name.strip()
    row = query("SELECT id FROM suppliers WHERE lower(name)=lower(?)", (name,))
    if len(row):
        return int(row.iloc[0]["id"])
    sid = execute("INSERT INTO suppliers(name,contact,created_at) VALUES(?,?,?)", (name, contact, now()))
    log("supplier", sid, "created", name)
    return sid


# ----------------------------------------------------------------------------
# DOCUMENT TEXT EXTRACTION
# ----------------------------------------------------------------------------
def extract_text(uploaded):
    name = uploaded.name.lower()
    data = uploaded.getvalue()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader
            return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(data)).pages)
        if name.endswith(".docx"):
            import docx
            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for t in d.tables:
                for r in t.rows:
                    parts.append(" | ".join(c.text for c in r.cells))
            return "\n".join(parts)
        return data.decode("utf-8", errors="ignore")
    except Exception as e:  # noqa
        st.error(f"Could not read {uploaded.name}: {e}")
        return ""


# ----------------------------------------------------------------------------
# EVALUATION ENGINES
# ----------------------------------------------------------------------------
def rubric_prompt():
    g = "\n".join(f"- {n}: {d}" for n, d, _ in GATES)
    c = "\n".join(f'- key="{k}" | {n} | weight {w}: {d}' for k, n, w, d, _ in CRITERIA)
    return f"""You are a strict, evidence-based procurement evaluator for an EV manufacturer.
Evaluate ONE supplier RFP response against this rubric (EV Traction Battery Supplier RFP template v1.0).
Unsupported claims receive no credit. Credit only traceable evidence (test reports, certifications, drawings, controlled plans).
Distinguish validated vs planned vs assumed evidence. Do not invent facts that are not in the response.

MANDATORY GATES (result must be exactly "Pass", "Fail" or "Clarify"):
{g}

WEIGHTED CRITERIA (score each 0-5; anchors 5 Leading, 4 Strong, 3 Acceptable, 2 Weak, 1 Poor, 0 No evidence):
{c}

Return ONLY valid JSON (no markdown) with this exact shape:
{{
 "gates": {{"<gate name>": {{"result": "Pass|Fail|Clarify", "evidence": "short quote/paraphrase", "gap": "what is missing"}}}},
 "scores": {{"<criterion key>": {{"score": 0-5 integer, "evidence": "...", "gaps": "..."}}}},
 "summary": "3-5 sentence overall assessment incl. key risks",
 "clarifications": ["questions to send the supplier"],
 "extracted": {{
   "quoted_cost": null or number, "currency": "USD" or null,
   "batteries": [{{"chemistry": "", "capacity_kwh": null, "voltage_v": null, "price_per_kwh": null}}],
   "materials": [{{"material": "", "category": "cathode|anode|electrolyte|structural|thermal|other", "origin": "", "recycled_pct": null}}]
 }}
}}
Use the gate names and criterion keys exactly as given."""


def parse_json(text):
    text = text.strip()
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.M).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        s, e = text.find("{"), text.rfind("}")
        return json.loads(text[s:e + 1])


def evaluate_llm(text, api_key, model):
    if anthropic is None:
        raise RuntimeError("`anthropic` package not installed. Run: pip install anthropic")
    client = anthropic.Anthropic(api_key=api_key)
    msg = client.messages.create(
        model=model, max_tokens=6000, system=rubric_prompt(),
        messages=[{"role": "user", "content": "SUPPLIER RESPONSE:\n\n" + text[:120000]}],
    )
    out = "".join(b.text for b in msg.content if getattr(b, "type", "") == "text")
    data = parse_json(out)
    data["_tokens"] = (msg.usage.input_tokens, msg.usage.output_tokens)
    return data


def evaluate_heuristic(text):
    """Offline keyword screen. Rough triage only - NOT a substitute for the LLM evaluation or human review."""
    low = text.lower()
    gates = {}
    for n, _, kws in GATES:
        hits = [k for k in kws if k in low]
        res = "Pass" if len(hits) >= 3 else "Clarify" if hits else "Fail"
        gates[n] = {"result": res, "evidence": f"Keywords found: {', '.join(hits) or 'none'}", "gap": "Keyword screen only - verify evidence."}
    scores = {}
    for k, n, w, d, kws in CRITERIA:
        hits = [x for x in kws if x in low]
        ratio = len(hits) / max(len(kws), 1)
        s = 0 if not hits else min(5, 1 + round(ratio * 4))
        scores[k] = {"score": s, "evidence": f"Keywords: {', '.join(hits) or 'none'}", "gaps": "Keyword screen only."}
    return {"gates": gates, "scores": scores,
            "summary": "Offline heuristic keyword screen (no API key used). Use for quick triage; run the AI evaluator for evidence-based scoring.",
            "clarifications": [], "extracted": {}, "_tokens": (0, 0)}


def run_evaluation(resp_id, mode, api_key, model):
    r = query("SELECT * FROM responses WHERE id=?", (resp_id,)).iloc[0]
    t0 = time.time()
    status, err, data = "ok", "", None
    try:
        data = evaluate_llm(r["text"], api_key, model) if mode == "AI (Claude API)" else evaluate_heuristic(r["text"])
    except Exception as e:  # noqa
        status, err = "error", str(e)[:500]
    dur = round(time.time() - t0, 2)
    if status == "error":
        rid = execute("""INSERT INTO runs(response_id,mode,model,started,duration_s,status,error,total_score,rating,gates_passed)
                         VALUES(?,?,?,?,?,?,?,?,?,?)""", (resp_id, mode, model, now(), dur, status, err, 0, "Error", 0))
        log("run", rid, "failed", err)
        return rid, None
    # normalise
    gates = {n: data.get("gates", {}).get(n, {"result": "Clarify", "evidence": "", "gap": "Not returned"}) for n, _, _ in GATES}
    for n in gates:
        if gates[n].get("result") not in ("Pass", "Fail", "Clarify"):
            gates[n]["result"] = "Clarify"
    scores = {}
    for k in CRIT_KEYS:
        s = data.get("scores", {}).get(k, {"score": 0, "evidence": "", "gaps": "Not returned"})
        s["score"] = max(0, min(5, int(s.get("score", 0) or 0)))
        scores[k] = s
    total = weighted_total(scores)
    ok = all(g["result"] == "Pass" for g in gates.values())
    rating = band(total, ok)
    summary = data.get("summary", "")
    if data.get("clarifications"):
        summary += "\n\nClarifications to request:\n- " + "\n- ".join(map(str, data["clarifications"]))
    tin, tout = data.get("_tokens", (0, 0))
    rid = execute("""INSERT INTO runs(response_id,mode,model,started,duration_s,total_score,rating,gates_passed,gates_json,
                     scores_json,summary,tokens_in,tokens_out,status,error) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                  (resp_id, mode, model if mode.startswith("AI") else "heuristic", now(), dur, total, rating, int(ok),
                   json.dumps(gates), json.dumps(scores), summary, tin, tout, "ok", ""))
    log("run", rid, "completed", f"response {resp_id}: {total}/100, {rating}")
    return rid, data.get("extracted", {})


def store_extracted(supplier_name, ext):
    n = 0
    for b in ext.get("batteries", []) or []:
        if b.get("chemistry"):
            execute("INSERT INTO supplier_batteries(supplier,chemistry,capacity_kwh,voltage_v,price_per_kwh,source) VALUES(?,?,?,?,?,?)",
                    (supplier_name, b.get("chemistry"), b.get("capacity_kwh"), b.get("voltage_v"), b.get("price_per_kwh"), "extracted"))
            n += 1
    for m in ext.get("materials", []) or []:
        if m.get("material"):
            execute("INSERT INTO materials(supplier,material,category,origin,recycled_pct,source) VALUES(?,?,?,?,?,?)",
                    (supplier_name, m.get("material"), m.get("category"), m.get("origin"), m.get("recycled_pct"), "extracted"))
            n += 1
    if n:
        log("catalog", supplier_name, "auto-added", f"{n} battery/material rows extracted from evaluation")
    return n


# ----------------------------------------------------------------------------
# UI HELPERS
# ----------------------------------------------------------------------------
SUV_SVG = """
<svg viewBox="0 0 320 130" width="260" xmlns="http://www.w3.org/2000/svg" role="img" aria-label="EV SUV">
  <ellipse cx="160" cy="118" rx="140" ry="6" fill="#00000022"/>
  <path d="M18 92 L18 68 Q20 56 40 53 L78 49 L102 26 Q106 21 114 21 L206 21 Q216 21 222 28 L246 49 L282 55 Q298 58 298 70 L298 92 Z" fill="#0b6e4f"/>
  <path d="M90 49 L108 31 L152 31 L152 49 Z" fill="#d6eef8"/>
  <path d="M160 31 L204 31 L225 49 L160 49 Z" fill="#d6eef8"/>
  <rect x="18" y="84" width="280" height="8" fill="#084c37"/>
  <rect x="282" y="62" width="16" height="7" rx="2" fill="#ffe27a"/>
  <rect x="18" y="62" width="12" height="7" rx="2" fill="#ff6b6b"/>
  <path d="M170 60 l-10 16 h8 l-4 14 l14 -20 h-9 l6 -10 z" fill="#ffe27a"/>
  <circle cx="78" cy="94" r="22" fill="#222"/><circle cx="78" cy="94" r="10" fill="#cfd8dc"/>
  <circle cx="240" cy="94" r="22" fill="#222"/><circle cx="240" cy="94" r="10" fill="#cfd8dc"/>
</svg>
"""


def rating_badge(r):
    return {"Preferred": "🟢 Preferred", "Competitive": "🔵 Competitive", "Conditional": "🟠 Conditional",
            "Do not progress": "🔴 Do not progress"}.get(r, "⛔ " + str(r))


def latest_view():
    return query("""
        SELECT r.id AS response_id, s.name AS supplier, r.rfp_ref, r.title, r.cost, r.currency, r.created_at,
               ru.id AS run_id, ru.total_score, ru.rating, ru.gates_passed
        FROM responses r JOIN suppliers s ON s.id=r.supplier_id
        LEFT JOIN runs ru ON ru.id=(SELECT MAX(id) FROM runs WHERE response_id=r.id AND status='ok')
        ORDER BY s.name""")


# ----------------------------------------------------------------------------
# PAGES
# ----------------------------------------------------------------------------
def page_dashboard():
    c1, c2 = st.columns([1, 3])
    with c1:
        st.markdown(SUV_SVG, unsafe_allow_html=True)
    with c2:
        st.title("EV Battery Supplier RFP Evaluator")
        st.caption("Compare and validate supplier responses against the EV Traction Battery Supplier RFP template (100-point matrix).")

    lv = latest_view()
    if lv.empty:
        st.info("No supplier responses yet. Go to **Import Response** to add one (or tick *Load demo data* in the sidebar).")
        return
    n_sup = lv["supplier"].nunique()
    n_resp = len(lv)
    n_rfp = lv["rfp_ref"].replace("", pd.NA).nunique()
    evaluated = lv.dropna(subset=["total_score"])
    k = st.columns(5)
    k[0].metric("Suppliers", n_sup)
    k[1].metric("Responses / RFPs", n_resp)
    k[2].metric("Distinct RFP refs", n_rfp)
    k[3].metric("Avg score /100", round(evaluated["total_score"].mean(), 1) if len(evaluated) else "-")
    k[4].metric("Total quoted cost", f"{lv['cost'].fillna(0).sum():,.0f}")

    st.subheader("Supplier scorecard")
    agg = lv.groupby("supplier").agg(
        rfps=("response_id", "count"),
        score=("total_score", "mean"),
        cost=("cost", "sum"),
        gates=("gates_passed", "min"),
    ).reset_index()
    agg["rating"] = [band(s, bool(g)) if pd.notna(s) else "Not evaluated" for s, g in zip(agg["score"], agg["gates"])]
    agg["stars (0-5)"] = agg["score"].apply(lambda s: stars(s) if pd.notna(s) else None)
    agg["rating"] = agg["rating"].apply(rating_badge)
    agg = agg.sort_values("score", ascending=False)
    show = agg.rename(columns={"supplier": "Supplier", "rfps": "Total RFPs", "score": "Score /100", "cost": "Cost involved",
                               "rating": "Rating"})[["Supplier", "Total RFPs", "Score /100", "stars (0-5)", "Rating", "Cost involved"]]
    st.dataframe(show.round(1), use_container_width=True, hide_index=True)
    if len(evaluated):
        c1, c2 = st.columns(2)
        c1.markdown("**Score by supplier**")
        c1.bar_chart(agg.set_index("supplier")["score"])
        c2.markdown("**Cost by supplier**")
        c2.bar_chart(agg.set_index("supplier")["cost"])
    with st.expander("All responses (latest evaluation per response)"):
        st.dataframe(lv.drop(columns=["response_id", "run_id"]), use_container_width=True, hide_index=True)


def page_import():
    st.header("Import supplier RFP response")
    with st.form("import", clear_on_submit=True):
        c1, c2 = st.columns(2)
        supplier = c1.text_input("Supplier name *")
        contact = c2.text_input("Contact (optional)")
        c3, c4, c5 = st.columns(3)
        rfp_ref = c3.text_input("RFP reference", placeholder="e.g. RFP-2026-BAT-01")
        cost = c4.number_input("Quoted cost / TCO", min_value=0.0, step=1000.0)
        currency = c5.selectbox("Currency", ["USD", "EUR", "INR", "GBP", "CNY", "JPY"])
        files = st.file_uploader("Response file(s): PDF, DOCX or TXT", type=["pdf", "docx", "txt", "md"], accept_multiple_files=True)
        pasted = st.text_area("...or paste response text", height=150)
        ok = st.form_submit_button("Import response")
    if ok:
        if not supplier.strip():
            st.error("Supplier name is required.")
            return
        text = pasted.strip()
        fname = "pasted"
        for f in files or []:
            text += "\n\n" + extract_text(f)
            fname = f.name if len(files) == 1 else f"{len(files)} files"
        if not text.strip():
            st.error("No text found. Upload a readable file or paste text.")
            return
        sid = get_or_create_supplier(supplier, contact)
        rid = execute("INSERT INTO responses(supplier_id,rfp_ref,title,text,cost,currency,filename,created_at) VALUES(?,?,?,?,?,?,?,?)",
                      (sid, rfp_ref, fname, text, cost, currency, fname, now()))
        log("response", rid, "imported", f"{supplier} | {rfp_ref} | {fname} | {len(text)} chars | cost {cost} {currency}")
        st.success(f"Imported response #{rid} for {supplier} ({len(text):,} characters). Go to **Run Evaluator**.")

    st.subheader("Imported responses")
    df = query("""SELECT r.id, s.name AS supplier, r.rfp_ref, r.filename, r.cost, r.currency, length(r.text) AS chars, r.created_at
                  FROM responses r JOIN suppliers s ON s.id=r.supplier_id ORDER BY r.id DESC""")
    st.dataframe(df, use_container_width=True, hide_index=True)
    if len(df):
        with st.expander("Edit cost or delete a response"):
            rid = st.selectbox("Response id", df["id"].tolist())
            row = df[df["id"] == rid].iloc[0]
            new_cost = st.number_input("Cost", value=float(row["cost"] or 0), step=1000.0, key="editcost")
            c1, c2 = st.columns(2)
            if c1.button("Save cost"):
                execute("UPDATE responses SET cost=? WHERE id=?", (new_cost, int(rid)))
                log("response", rid, "cost changed", f"{row['cost']} -> {new_cost}")
                st.rerun()
            if c2.button("Delete response", type="secondary"):
                execute("DELETE FROM runs WHERE response_id=?", (int(rid),))
                execute("DELETE FROM responses WHERE id=?", (int(rid),))
                log("response", rid, "deleted", f"{row['supplier']} {row['rfp_ref']}")
                st.rerun()


def page_run(api_key, model):
    st.header("Run evaluator")
    with st.expander("Validation reference: RFP template rubric", expanded=False):
        st.markdown("**Mandatory gates** (all must pass before ranking)")
        st.dataframe(pd.DataFrame([(n, d) for n, d, _ in GATES], columns=["Gate", "Minimum submission"]), hide_index=True, use_container_width=True)
        st.markdown("**Weighted criteria** (score 0-5; points = score/5 x weight)")
        st.dataframe(pd.DataFrame([(n, w, d) for _, n, w, d, _ in CRITERIA], columns=["Criterion", "Weight", "Evidence for a high score"]),
                     hide_index=True, use_container_width=True)
        st.markdown("Thresholds: 85-100 Preferred | 75-84 Competitive | 65-74 Conditional | <65 Do not progress | any failed gate: not ranked")

    resp = query("SELECT r.id, s.name AS supplier, r.rfp_ref, r.filename FROM responses r JOIN suppliers s ON s.id=r.supplier_id ORDER BY r.id DESC")
    if resp.empty:
        st.info("Import at least one supplier response first.")
        return
    labels = {int(r.id): f"#{r.id} - {r.supplier} - {r.rfp_ref or 'no ref'} ({r.filename})" for r in resp.itertuples()}
    chosen = st.multiselect("Responses to evaluate", list(labels), format_func=labels.get, default=list(labels)[:1])
    mode = st.radio("Evaluation mode", ["AI (Claude API)", "Offline heuristic (no key, rough triage)"], horizontal=True)
    auto = st.checkbox("Auto-add extracted batteries / materials to the Batteries tab", value=True)
    if mode.startswith("AI") and not api_key:
        st.warning("Enter your Anthropic API key in the sidebar to use AI mode.")
    if st.button("Run evaluator", type="primary", disabled=not chosen or (mode.startswith("AI") and not api_key)):
        bar = st.progress(0.0)
        for i, rid in enumerate(chosen, 1):
            run_id, ext = run_evaluation(rid, "AI (Claude API)" if mode.startswith("AI") else "Heuristic", api_key, model)
            r = query("SELECT * FROM runs WHERE id=?", (run_id,)).iloc[0]
            if r["status"] == "error":
                st.error(f"{labels[rid]}: {r['error']}")
            else:
                st.success(f"{labels[rid]} -> **{r['total_score']}/100** | {rating_badge(r['rating'])}")
                if auto and ext:
                    sname = resp[resp["id"] == rid].iloc[0]["supplier"]
                    store_extracted(sname, ext)
            bar.progress(i / len(chosen))
        st.info("Open **Runs & Performance** for full gate results and evidence.")


def page_runs():
    st.header("Evaluator runs & performance")
    runs = query("""SELECT ru.id AS run, s.name AS supplier, r.rfp_ref, ru.mode, ru.model, ru.started, ru.duration_s,
                    ru.total_score, ru.rating, ru.gates_passed, ru.tokens_in, ru.tokens_out, ru.status
                    FROM runs ru JOIN responses r ON r.id=ru.response_id JOIN suppliers s ON s.id=r.supplier_id ORDER BY ru.id DESC""")
    if runs.empty:
        st.info("No runs yet.")
        return
    ok = runs[runs["status"] == "ok"]
    k = st.columns(5)
    k[0].metric("Total runs", len(runs))
    k[1].metric("Successful", len(ok))
    k[2].metric("Failed", len(runs) - len(ok))
    k[3].metric("Avg duration (s)", round(ok["duration_s"].mean(), 1) if len(ok) else "-")
    k[4].metric("Tokens used", int(ok["tokens_in"].fillna(0).sum() + ok["tokens_out"].fillna(0).sum()))
    st.dataframe(runs, use_container_width=True, hide_index=True)

    if len(ok):
        st.subheader("Score trend per supplier (across runs)")
        piv = ok.sort_values("run").pivot_table(index="run", columns="supplier", values="total_score", aggfunc="last").ffill()
        st.line_chart(piv)

    st.subheader("Run detail")
    rid = st.selectbox("Select run", runs["run"].tolist())
    r = query("SELECT * FROM runs WHERE id=?", (int(rid),)).iloc[0]
    if r["status"] == "error":
        st.error(r["error"])
        return
    st.markdown(f"### {r['total_score']}/100 - {rating_badge(r['rating'])}  \n_{r['mode']} | {r['model']} | {r['started']}_")
    gates = json.loads(r["gates_json"])
    scores = json.loads(r["scores_json"])
    st.markdown("**Mandatory gates**")
    st.dataframe(pd.DataFrame([(n, g["result"], g.get("evidence", ""), g.get("gap", "")) for n, g in gates.items()],
                              columns=["Gate", "Result", "Evidence", "Gap"]), hide_index=True, use_container_width=True)
    rows = []
    for k_, n, w, _, _ in CRITERIA:
        s = scores[k_]["score"]
        rows.append((n, w, s, SCORE_ANCHORS[s], round(s / 5 * w, 1), scores[k_].get("evidence", ""), scores[k_].get("gaps", "")))
    sdf = pd.DataFrame(rows, columns=["Criterion", "Weight", "Score 0-5", "Anchor", "Points", "Evidence", "Gaps"])
    st.markdown("**Weighted criteria**")
    st.dataframe(sdf, hide_index=True, use_container_width=True)
    st.bar_chart(sdf.set_index("Criterion")[["Points", "Weight"]])
    st.markdown("**Summary**")
    st.write(r["summary"])
    st.download_button("Download run (JSON)", json.dumps({"run": int(rid), "gates": gates, "scores": scores, "summary": r["summary"]}, indent=2),
                       file_name=f"run_{rid}.json")

    # change vs previous run for the same response
    prev = query("""SELECT id,total_score FROM runs WHERE status='ok' AND response_id=? AND id<?
                    ORDER BY id DESC LIMIT 1""", (int(r["response_id"]), int(rid)))
    if len(prev):
        d = round(r["total_score"] - prev.iloc[0]["total_score"], 1)
        st.info(f"Change vs previous run #{int(prev.iloc[0]['id'])} of this response: {d:+} points")


def editable_table(table, cols, title, defaults=None):
    df = query(f"SELECT {', '.join(cols)} FROM {table}")
    edited = st.data_editor(df, num_rows="dynamic", use_container_width=True, hide_index=True, key=f"ed_{table}")
    if st.button(f"Save {title}", key=f"sv_{table}"):
        before = df.to_json(orient="records")
        execute(f"DELETE FROM {table}")
        with conn() as c:
            for _, row in edited.iterrows():
                if all(pd.isna(v) or v == "" for v in row.values):
                    continue
                vals = [None if pd.isna(v) else v for v in row.values]
                c.execute(f"INSERT INTO {table}({', '.join(cols)}) VALUES({', '.join('?' * len(cols))})", vals)
        log(table, "-", "table saved", f"before: {before[:1500]} || after: {edited.to_json(orient='records')[:1500]}")
        st.success("Saved and logged in change history.")
        st.rerun()


def page_batteries():
    st.header("EV batteries & battery materials")
    st.markdown("**1. Battery types our company supports**")
    editable_table("company_batteries", ["chemistry", "application", "notes"], "company batteries")
    st.markdown("**2. Batteries provided by suppliers**")
    editable_table("supplier_batteries", ["supplier", "chemistry", "capacity_kwh", "voltage_v", "price_per_kwh", "source"], "supplier batteries")
    st.markdown("**3. Battery materials from suppliers**")
    editable_table("materials", ["supplier", "material", "category", "origin", "recycled_pct", "source"], "materials")

    comp = query("SELECT DISTINCT lower(chemistry) AS c FROM company_batteries")["c"].tolist()
    sb = query("SELECT supplier, chemistry, capacity_kwh, price_per_kwh FROM supplier_batteries")
    if comp and len(sb):
        sb["supported by us?"] = sb["chemistry"].str.lower().apply(lambda x: "Yes" if any(c in x or x in c for c in comp) else "No")
        st.markdown("**Fit check: supplier batteries vs our supported chemistries**")
        st.dataframe(sb, hide_index=True, use_container_width=True)


def page_history():
    st.header("Memory & history of changes")
    st.caption("Everything imported, edited, deleted or run is stored in a local SQLite file and survives restarts.")
    a = query("SELECT * FROM audit ORDER BY id DESC")
    ent = st.multiselect("Filter by entity", sorted(a["entity"].unique().tolist()) if len(a) else [])
    if ent:
        a = a[a["entity"].isin(ent)]
    st.dataframe(a, use_container_width=True, hide_index=True)
    if len(a):
        st.download_button("Export history (CSV)", a.to_csv(index=False), file_name="rfp_audit_history.csv")
    st.caption(f"Database file: {os.path.abspath(DB_PATH)}")


def page_suppliers():
    st.header("Supplier history & expertise")
    sup = query("SELECT name FROM suppliers ORDER BY name")["name"].tolist()
    if not sup:
        st.info("No suppliers yet. Import a response first.")
        return

    # --- overview of all suppliers
    prof = query("SELECT * FROM supplier_profiles")
    perf = query("""SELECT s.name AS supplier, COUNT(DISTINCT r.id) AS rfps, COUNT(ru.id) AS runs, ROUND(AVG(ru.total_score),1) AS avg_score,
                    ROUND(MAX(ru.total_score),1) AS best_score, ROUND((SELECT SUM(cost) FROM responses WHERE supplier_id=s.id),0) AS total_cost
                    FROM suppliers s LEFT JOIN responses r ON r.supplier_id=s.id
                    LEFT JOIN runs ru ON ru.response_id=r.id AND ru.status='ok' GROUP BY s.name""")
    ov = perf.merge(prof[["supplier", "years_in_business", "specialties", "expertise_level"]], on="supplier", how="left")
    ov = ov.rename(columns={"expertise_level": "expertise (1-5)", "years_in_business": "years"})
    st.dataframe(ov, use_container_width=True, hide_index=True)

    name = st.selectbox("Supplier", sup)
    p = prof[prof["supplier"] == name]
    p = p.iloc[0].to_dict() if len(p) else {}
    t1, t2, t3 = st.tabs(["Expertise profile", "Evaluation history", "Change log"])

    with t1:
        with st.form(f"prof_{name}"):
            c1, c2, c3 = st.columns(3)
            years = c1.number_input("Years in business", 0, 200, int(p.get("years_in_business") or 0))
            hq = c2.text_input("HQ / manufacturing sites", p.get("hq") or "")
            lvl = c3.slider("Expertise in EV batteries (buyer view)", 1, 5, int(p.get("expertise_level") or 3))
            spec = st.text_input("Specialties (e.g. LFP cells, pack integration, BMS, thermal)", p.get("specialties") or "")
            cert = st.text_input("Certifications (IATF 16949, ISO 9001, UL 2580...)", p.get("certifications") or "")
            cust = st.text_input("Reference customers / programs", p.get("key_customers") or "")
            notes = st.text_area("Notes", p.get("notes") or "")
            if st.form_submit_button("Save profile"):
                new = dict(years_in_business=years, hq=hq, specialties=spec, certifications=cert, key_customers=cust,
                           expertise_level=lvl, notes=notes)
                execute("""INSERT INTO supplier_profiles(supplier,years_in_business,hq,specialties,certifications,key_customers,
                           expertise_level,notes,updated) VALUES(?,?,?,?,?,?,?,?,?)
                           ON CONFLICT(supplier) DO UPDATE SET years_in_business=excluded.years_in_business, hq=excluded.hq,
                           specialties=excluded.specialties, certifications=excluded.certifications, key_customers=excluded.key_customers,
                           expertise_level=excluded.expertise_level, notes=excluded.notes, updated=excluded.updated""",
                        (name, years, hq, spec, cert, cust, lvl, notes, now()))
                changed = {k: f"{p.get(k)} -> {v}" for k, v in new.items() if str(p.get(k) or "") != str(v or "")}
                log("supplier_profile", name, "saved", json.dumps(changed))
                st.success("Profile saved and logged.")
                st.rerun()

        st.markdown("**Evidence-based expertise (from evaluator runs)**")
        rs = query("""SELECT ru.scores_json FROM runs ru JOIN responses r ON r.id=ru.response_id JOIN suppliers s ON s.id=r.supplier_id
                      WHERE s.name=? AND ru.status='ok' AND ru.id=(SELECT MAX(id) FROM runs WHERE response_id=r.id AND status='ok')""", (name,))
        if rs.empty:
            st.caption("No evaluated responses yet for this supplier.")
        else:
            rows = []
            for k_, n, w, _, _ in CRITERIA:
                vals = [json.loads(j)[k_]["score"] for j in rs["scores_json"]]
                rows.append((n, round(sum(vals) / len(vals), 1), w))
            ex = pd.DataFrame(rows, columns=["Criterion", "Avg score 0-5", "Weight"])
            st.bar_chart(ex.set_index("Criterion")["Avg score 0-5"])
            top = ex.sort_values("Avg score 0-5", ascending=False)
            c1, c2 = st.columns(2)
            c1.success("Strongest: " + "; ".join(top.head(3)["Criterion"]))
            c2.warning("Weakest: " + "; ".join(top.tail(3)["Criterion"]))

        b = query("SELECT chemistry, capacity_kwh, voltage_v, price_per_kwh FROM supplier_batteries WHERE supplier=?", (name,))
        m = query("SELECT material, category, origin, recycled_pct FROM materials WHERE supplier=?", (name,))
        c1, c2 = st.columns(2)
        c1.markdown("**Batteries offered**")
        c1.dataframe(b, hide_index=True, use_container_width=True)
        c2.markdown("**Materials supplied**")
        c2.dataframe(m, hide_index=True, use_container_width=True)

    with t2:
        h = query("""SELECT r.id AS response, r.rfp_ref, r.cost, r.currency, ru.id AS run, ru.started, ru.mode, ru.total_score, ru.rating, ru.gates_passed
                     FROM responses r JOIN suppliers s ON s.id=r.supplier_id LEFT JOIN runs ru ON ru.response_id=r.id AND ru.status='ok'
                     WHERE s.name=? ORDER BY ru.id""", (name,))
        st.dataframe(h, hide_index=True, use_container_width=True)
        hh = h.dropna(subset=["total_score"])
        if len(hh) > 1:
            st.markdown("**Score over time**")
            st.line_chart(hh.set_index("started")["total_score"])
        elif len(hh) == 1:
            st.caption("Only one evaluation so far; trend appears after more runs.")

    with t3:
        a = query("SELECT ts, entity, action, detail FROM audit WHERE (entity='supplier_profile' AND entity_id=?) OR detail LIKE ? ORDER BY id DESC",
                  (name, f"%{name}%"))
        st.dataframe(a, hide_index=True, use_container_width=True)


def load_demo():
    demos = {
        "VoltCore Cells": ("RFP-DEMO-01", 48_000_000, """VoltCore proposes an NMC811 pack, 90 kWh gross, 400V. We hold IATF 16949 and ISO 9001 certificates.
Compliance plan covers UN ECE R100 and UL 2580 with a deviations register and named accountable owner.
Sealing: IP67 enclosure, helium leak test and pressure decay at end-of-line; coolant containment and venting concept included.
Thermal: thermal runaway propagation test per abuse protocol, gas detection and mitigation; cooling plate model with temperature uniformity data and sensor placement.
Vibration, shock, road-load durability and thermal cycling reports with post-test insulation resistance are attached.
Traceability: cell-to-pack genealogy with serial and batch records; recall support. Cybersecurity: ISO/SAE 21434 plan, secure boot, SBOM, penetration test, incident response.
Materials: NMC cathode, graphite anode, LiPF6 electrolyte; cobalt and nickel due diligence, IMDS declarations, carbon footprint roadmap, 12% recycled content.
Fast-charge 10-80% in 22 minutes, cycle life 2000, calendar life 15 years, 8 year warranty, SOH estimation. Take-back, recycling partner and second-life reuse plan. BCP, audited financials, capacity 20 GWh, PPAP/APQP plan."""),
        "GreenAmp Energy": ("RFP-DEMO-01", 41_500_000, """GreenAmp offers an LFP pack, 75 kWh, 350V. ISO 9001 certified; IATF transition plan in progress.
Sealed enclosure with IP67 target. Thermal management with liquid cooling. Vibration testing planned.
We will provide recycling through a partner. Price competitive. Capacity 8 GWh."""),
    }
    for name, (ref, cost, txt) in demos.items():
        sid = get_or_create_supplier(name)
        rid = execute("INSERT INTO responses(supplier_id,rfp_ref,title,text,cost,currency,filename,created_at) VALUES(?,?,?,?,?,?,?,?)",
                      (sid, ref, "demo", txt, cost, "USD", "demo.txt", now()))
        log("response", rid, "imported", f"DEMO {name}")
    execute("INSERT INTO company_batteries(chemistry,application,notes) VALUES('NMC','SUV long range','demo')")
    execute("INSERT INTO company_batteries(chemistry,application,notes) VALUES('LFP','Entry trims','demo')")


# ----------------------------------------------------------------------------
# MAIN
# ----------------------------------------------------------------------------
def main():
    st.set_page_config(page_title="EV RFP Evaluator", page_icon="🚙", layout="wide")
    init_db()

    with st.sidebar:
        st.markdown(SUV_SVG, unsafe_allow_html=True)
        st.subheader("🔑 API key")
        key = st.text_input("Anthropic API key", type="password", value=st.session_state.get("api_key", os.environ.get("ANTHROPIC_API_KEY", "")),
                            help="Held in this browser session only. It is never written to the database.")
        st.session_state["api_key"] = key
        model = st.text_input("Model", value=st.session_state.get("model", "claude-sonnet-5-5"))
        st.session_state["model"] = model
        st.caption("Key status: " + ("✅ entered" if key else "❌ missing (offline heuristic mode still works)"))
        st.divider()
        if st.button("Load demo data"):
            load_demo()
            st.rerun()

    tabs = st.tabs(["📊 Dashboard", "📥 Import Response", "▶️ Run Evaluator", "📈 Runs & Performance",
                    "🏭 Suppliers", "🔋 Batteries & Materials", "🕘 History & Memory"])
    with tabs[0]:
        page_dashboard()
    with tabs[1]:
        page_import()
    with tabs[2]:
        page_run(key, model)
    with tabs[3]:
        page_runs()
    with tabs[4]:
        page_suppliers()
    with tabs[5]:
        page_batteries()
    with tabs[6]:
        page_history()


if __name__ == "__main__":
    main()
